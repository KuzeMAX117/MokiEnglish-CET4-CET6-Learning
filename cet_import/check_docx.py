# -*- coding: utf-8 -*-
"""Check docx structure including tables"""
from docx import Document
import re
import os

cache_dir = r"D:\202401\study\Hao-Topic\cet_import\docx_cache"

# Check the first file (CET6 set 1 2015)
fname = sorted(os.listdir(cache_dir))[0]
fpath = os.path.join(cache_dir, fname)
doc = Document(fpath)

print(f"File: {fname}")
print(f"Paragraphs: {len(doc.paragraphs)}")
print(f"Tables: {len(doc.tables)}")

# Check tables
for i, table in enumerate(doc.tables):
    print(f"\nTable {i}: {len(table.rows)} rows x {len(table.columns)} cols")
    for row in table.rows[:3]:
        cells = [cell.text[:30] for cell in row.cells]
        print(f"  {cells}")

# Check all text including tables
all_text = []
for p in doc.paragraphs:
    all_text.append(p.text)
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            all_text.append(cell.text)

full_text = "\n".join(all_text)

# Search for answer patterns
print("\n--- Searching for answers ---")
range_matches = list(re.finditer(r"(\d+)[-–](\d+)\s*[：:]\s*([A-D]{1,20})", full_text))
print(f"Range answer matches: {len(range_matches)}")
for m in range_matches:
    print(f"  {m.group()}")

# Check for answer section markers
for marker in ["参考答案", "答案", "Answer Key", "answer key", "ANSWER"]:
    idx = full_text.rfind(marker)
    if idx >= 0:
        print(f"Found '{marker}' at pos {idx}")
        print(f"  Context: {full_text[idx:idx+200]}")
