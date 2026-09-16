# -*- coding: utf-8 -*-
"""Check CET4 set 1 file for answers"""
from docx import Document
import re
import os

cache_dir = r"D:\202401\study\Topic\cet_import\docx_cache"

# Check file index 3 (CET4 set 1 2015)
fname = sorted(os.listdir(cache_dir))[3]
fpath = os.path.join(cache_dir, fname)
doc = Document(fpath)

print(f"File: {fname}")
print(f"Paragraphs: {len(doc.paragraphs)}")
print(f"Tables: {len(doc.tables)}")

# Get all text
all_text = "\n".join([p.text for p in doc.paragraphs])
print(f"Total text length: {len(all_text)}")

# Search for answer patterns
range_matches = re.findall(r"(\d+)[-–](\d+)\s*[：:]\s*([A-D]{1,20})", all_text)
print(f"Range answers found: {range_matches}")

# Check for 参考答案
idx = all_text.rfind("参考答案")
if idx >= 0:
    print(f"\nFound 参考答案 at {idx}")
    print(f"Context: ...{all_text[max(0,idx-100):idx+300]}...")

# Show last 500 chars
print(f"\nLast 500 chars of text:")
print(all_text[-500:])
